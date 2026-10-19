"""Aucune formule ne part dans un export (audit de sécurité du 19 octobre 2026)."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.partage.tableur import forcer_le_texte, ligne_sure, texte_sur


def test_une_valeur_qui_ressemble_a_une_formule_est_desarmee_en_csv():
    assert texte_sur('=HYPERLINK("http://x","clic")') == '\'=HYPERLINK("http://x","clic")'
    for debut in ("+", "-", "@", "\t", "\r"):
        assert texte_sur(f"{debut}1+1").startswith("'")
    assert ligne_sure(["Rose", 12, None, "=1+1"]) == ["Rose", 12, None, "'=1+1"]


def test_une_valeur_ordinaire_ne_change_pas():
    assert texte_sur("ROSE COUTURE SARL") == "ROSE COUTURE SARL"
    assert texte_sur(-45000) == -45000


def test_la_cellule_excel_reste_du_texte():
    classeur = Workbook()
    feuille = classeur.active
    feuille.append(["01/10/2026", '=HYPERLINK("http://x","clic")', 1000.0])
    forcer_le_texte(feuille, feuille.max_row, range(1, 3))
    tampon = BytesIO()
    classeur.save(tampon)

    relu = load_workbook(BytesIO(tampon.getvalue())).active
    assert relu.cell(row=1, column=2).data_type == "s"
    assert relu.cell(row=1, column=3).data_type == "n"
