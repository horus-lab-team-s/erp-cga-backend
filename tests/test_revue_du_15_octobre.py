"""Les corrections de la revue qualité du 15 octobre 2026 (voir `Docs/revue-qualite-2026-10-15.md`).

Chaque cas garde un défaut trouvé par la revue ; les cas par les routes vivent à côté
de leur sujet (`test_honoraires.py`, `test_fonds_de_tiers.py`, `test_documents_de_la_cliente.py`).
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.contextes.comptabilite.adaptateurs.sortant.balance_auxiliaire_excel import (
    EnteteBalanceAuxiliaire,
    FicheDuTiers,
    rendre_la_balance_auxiliaire,
)
from app.contextes.comptabilite.domaine.balance_auxiliaire import (
    LigneBalanceAuxiliaire,
    NatureAuxiliaire,
)
from app.partage.en_tetes import disposition_en_ligne
from app.partage.horloge import a_douala


class TestLHeureDeDouala:
    def test_utc_plus_une_heure_toute_l_annee(self):
        assert a_douala(datetime(2026, 12, 31, 23, 30)) == datetime(2027, 1, 1, 0, 30)
        assert a_douala(datetime(2026, 7, 1, 9, 0)) == datetime(2026, 7, 1, 10, 0)


class TestLeNomDeFichier:
    def test_un_nom_accentue_ou_guillemete_ne_casse_pas_l_en_tete(self):
        en_tete = disposition_en_ligne('Bail d’Akwa "signé".pdf')

        en_tete.encode("latin-1")  # sinon le serveur lève, et le client reçoit une 500
        assert 'filename="Bail dAkwa signe.pdf"' in en_tete
        assert "filename*=UTF-8''Bail%20d%E2%80%99Akwa%20%22sign%C3%A9%22.pdf" in en_tete


class TestLExcelDesBalances:
    def _classeur(self, lignes, fiches):
        octets = rendre_la_balance_auxiliaire(
            lignes,
            EnteteBalanceAuxiliaire(
                niu="M081234567890P",
                denomination="SARL BATIMENT PLUS",
                exercice="2026",
                nature=NatureAuxiliaire.FOURNISSEURS,
                editee_le=datetime(2026, 10, 15, 11, 0),
            ),
            fiches,
        )
        return load_workbook(BytesIO(octets))

    def test_une_raison_sociale_en_egal_reste_du_texte(self):
        piege = '=HYPERLINK("http://exemple.invalid")'
        classeur = self._classeur(
            [LigneBalanceAuxiliaire("=1+1", Decimal(0), Decimal(0), Decimal(100))],
            {"=1+1": FicheDuTiers(denomination=piege, niu=None)},
        )
        cellules = classeur["Balance"][2]

        assert cellules[0].data_type == "s" and cellules[0].value == "=1+1"
        assert cellules[2].data_type == "s" and cellules[2].value == piege

    def test_les_montants_s_arrondissent_et_le_total_tombe_juste(self):
        lignes = [
            LigneBalanceAuxiliaire("A", Decimal(0), Decimal(0), Decimal("100.5")),
            LigneBalanceAuxiliaire("B", Decimal(0), Decimal(0), Decimal("200.5")),
        ]
        rangees = list(self._classeur(lignes, {})["Balance"].iter_rows(values_only=True))

        assert [r[5] for r in rangees[1:3]] == [101, 201]  # arrondi, pas tronqué
        assert rangees[-1][5] == 302  # la somme de ce que la colonne affiche

    def test_l_heure_d_edition_est_celle_de_douala(self):
        lire = self._classeur([], {})["À lire"]
        texte = " ".join(str(r[0]) for r in lire.iter_rows(values_only=True))

        assert "15/10/2026 à 12:00 (heure de Douala)" in texte


def test_la_date_du_jour_est_celle_de_douala():
    """16 octobre 2026 : à 00 h 30 à Douala (23 h 30 UTC), c'est déjà le lendemain. La
    suite, qui avait franchi minuit, a vu un RCCM « obtenu aujourd'hui » refusé comme
    futur : l'application prenait la date UTC."""
    from app.partage.horloge import aujourd_hui, horloge_figee

    with horloge_figee(datetime(2026, 10, 15, 23, 30)):
        assert aujourd_hui() == date(2026, 10, 16)


def test_le_jour_d_un_fait_horodate_est_celui_de_douala():
    """16 octobre 2026 : une déclaration déposée le 15 mai à 0 h 30 à Douala (le 14 à
    23 h 30 UTC) était datée du 14. Pour le retard, c'est le jour de Douala qui compte."""
    from app.partage.formats import moment_long
    from app.partage.horloge import jour_a_douala

    assert jour_a_douala(datetime(2026, 5, 14, 23, 30)) == date(2026, 5, 15)
    assert jour_a_douala(datetime(2026, 5, 14, 22, 59)) == date(2026, 5, 14)
    # Un lien qui expire à 23 h 30 UTC expire à 0 h 30 le lendemain, pour le lecteur.
    assert moment_long(datetime(2026, 8, 30, 23, 30)) == "31 août 2026 à 00h30"


def test_aucun_jour_n_est_tire_d_un_horodatage_utc():
    """Le garde-fou : `instant.date()` rend le jour UTC. Toute date d'un fait passe par
    `jour_a_douala` ou `aujourd_hui`. Seules les dates lues dans un texte (`strptime`)
    sont déjà locales."""
    import re
    from pathlib import Path

    motif = re.compile(r"(?<![\w.)\]])[A-Za-z_]\w*(?:\.[A-Za-z_]\w*)*\.date\(\)")
    fautes = []
    for fichier in Path("app").rglob("*.py"):
        if fichier.as_posix() == "app/partage/horloge.py":
            continue
        for numero, ligne in enumerate(fichier.read_text().splitlines(), 1):
            if "strptime" in ligne or ligne.lstrip().startswith(("#", ")")):
                continue
            if motif.search(ligne) and "a_douala(" not in ligne:
                fautes.append(f"{fichier}:{numero}: {ligne.strip()}")
    trouve = bool(fautes)
    assert not trouve, "\n".join(fautes)
