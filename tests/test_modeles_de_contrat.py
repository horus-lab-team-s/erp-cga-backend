"""Les modèles de contrat et le projet en Word (18 octobre 2026, point 23).

⚠️ CE QUI EST GARDÉ
* chaque modèle porte un canevas de cahier et des articles, sans citer d'article de loi ;
* partir d'un modèle ne demande pas l'agent ; le cahier déjà écrit n'est pas écrasé ;
  un modèle ne se verse qu'une fois, et seulement dans une rédaction ;
* un article retenu avec « [À COMPLÉTER » ne se soumet pas : la mention partirait au client ;
* le projet part en Word, fait des seuls articles retenus, numérotés, sans nom interne.
"""

from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO

import pytest

from app.contextes.souscription.adaptateurs.sortant.document_mission_word import (
    rendre_le_contrat_en_word,
)
from app.contextes.souscription.domaine.missions import (
    DecisionProposition,
    Mission,
    OrigineProposition,
    TransitionMissionRefusee,
    TypeMission,
)
from app.contextes.souscription.domaine.modeles_de_contrat import MODELES_DE_CONTRAT
from tests.conftest import exige_postgresql, ouvrir_une_session

LE = datetime(2026, 10, 18, 9, 0)
JURISTE = "p.moukouri@cga-brcg.cm"


def _texte_word(octets: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(octets))
    entete = " ".join(p.text for p in document.sections[0].header.paragraphs)
    return entete + "\n" + "\n".join(p.text for p in document.paragraphs)


def _redaction() -> Mission:
    return Mission(
        identifiant="MIS-1",
        type=TypeMission.REDACTION_CONTRAT,
        objet="Prestation",
        client="ETS X",
        responsable="C-007",
        ouverte_par="C-007",
        ouverte_le=LE,
    ).ecrire_le_cahier(type_de_contrat="contrat de prestation de services", cahier="x" * 30)


class TestLesModeles:
    def test_chaque_modele_a_son_cahier_et_ses_articles_sans_article_de_loi(self):
        assert len({m.code for m in MODELES_DE_CONTRAT}) == len(MODELES_DE_CONTRAT)
        for modele in MODELES_DE_CONTRAT:
            assert modele.cahier.count(":") >= 4 and len(modele.articles) >= 4
            for article in modele.articles:
                assert not re.search(r"\barticle\s+\d", article.texte, re.I), article.titre
                assert len(article.texte) <= 4000


class TestLaMentionACompleter:
    def test_retenue_telle_quelle_elle_bloque_la_soumission(self):
        m = _redaction().ajouter(
            [{"rubrique": "ARTICLE", "titre": "Prix", "texte": "[À COMPLÉTER : montant] FCFA."}],
            origine=OrigineProposition.MODELE,
        )
        retenue = m.decider(1, decision=DecisionProposition.RETENUE, par="C-007", a_l_instant=LE)
        with pytest.raises(TransitionMissionRefusee, match="À COMPLÉTER"):
            retenue.soumettre(par="C-007", a_l_instant=LE)

        reformulee = m.decider(
            1,
            decision=DecisionProposition.MODIFIEE,
            par="C-007",
            a_l_instant=LE,
            texte_retenu="Le prix est de 150 000 FCFA par mois.",
        )
        assert reformulee.soumettre(par="C-007", a_l_instant=LE).etat.value == "A_VALIDER"


class TestLeWord:
    def test_les_seuls_articles_retenus_numerotes_et_marques_projet(self):
        m = _redaction().ajouter(
            [
                {"rubrique": "ARTICLE", "titre": "Objet", "texte": "Tenue de la comptabilité."},
                {"rubrique": "ARTICLE", "titre": "Exclusivité", "texte": "Clause écartée."},
                {"rubrique": "ARTICLE", "titre": "Prix", "texte": "Brouillon."},
            ],
            origine=OrigineProposition.COLLABORATEUR,
        )
        m = m.decider(1, decision=DecisionProposition.RETENUE, par="C-007", a_l_instant=LE)
        m = m.decider(
            2, decision=DecisionProposition.ECARTEE, par="C-007", a_l_instant=LE, motif="non"
        )
        m = m.decider(
            3,
            decision=DecisionProposition.MODIFIEE,
            par="C-007",
            a_l_instant=LE,
            texte_retenu="150 000 FCFA par mois.\nPayables le 5.",
        )
        texte = _texte_word(rendre_le_contrat_en_word(m))

        assert "PROJET DE CONTRAT DE PRESTATION DE SERVICES" in texte
        assert "Article 1 · Objet" in texte and "Article 2 · Prix" in texte
        assert "Payables le 5." in texte and "Exclusivité" not in texte
        assert "PROJET · non validé" in texte and "C-007" not in texte


@exige_postgresql
class TestParLesRoutes:
    def _ouvrir(self, client, type_="REDACTION_CONTRAT") -> str:
        ouvrir_une_session(client, JURISTE)
        ouverte = client.post(
            "/missions", json={"type": type_, "objet": "Bail du dépôt", "client": "ETS X"}
        )
        assert ouverte.status_code == 201, ouverte.text
        return ouverte.json()["identifiant"]

    def test_partir_d_un_modele_sans_agent_puis_le_word(self, plateforme):
        client = plateforme
        identifiant = self._ouvrir(client)
        modeles = client.get("/missions/modeles-de-contrat").json()
        bail = next(m for m in modeles if m["code"] == "BAIL_COMMERCIAL")

        versee = client.post(f"/missions/{identifiant}/modele", json={"modele": "BAIL_COMMERCIAL"})
        assert versee.status_code == 200, versee.text
        mission = versee.json()
        assert mission["type_de_contrat"] == bail["libelle"] and mission["cahier"] == bail["cahier"]
        assert len(mission["propositions"]) == bail["articles"]
        assert {p["origine"] for p in mission["propositions"]} == {"MODELE"}
        assert {p["decision"] for p in mission["propositions"]} == {"EN_ATTENTE"}

        deux_fois = client.post(f"/missions/{identifiant}/modele", json={"modele": "TRAVAIL_CDI"})
        assert deux_fois.status_code == 409

        client.post(
            f"/missions/{identifiant}/propositions/1/decision",
            json={
                "decision": "MODIFIEE",
                "texte_retenu": "Le Bailleur donne à bail un entrepôt à Bonabéri.",
            },
        )
        word = client.get(f"/missions/{identifiant}/livrable/word")
        assert word.status_code == 200 and word.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml"
        )
        assert "entrepôt à Bonabéri" in _texte_word(word.content)

    def test_le_cahier_ecrit_est_garde(self, plateforme):
        client = plateforme
        identifiant = self._ouvrir(client)
        cahier = "Bail d'un entrepôt de 400 m² à Bonabéri, loyer convenu avec le client."
        client.post(
            f"/missions/{identifiant}/cahier",
            json={"type_de_contrat": "bail d'entrepôt", "cahier": cahier},
        )
        mission = client.post(
            f"/missions/{identifiant}/modele", json={"modele": "BAIL_COMMERCIAL"}
        ).json()
        assert mission["cahier"] == cahier and mission["type_de_contrat"] == "bail d'entrepôt"

    def test_ni_dans_une_analyse_ni_un_modele_inconnu(self, plateforme):
        client = plateforme
        analyse = self._ouvrir(client, "ANALYSE_CONTRAT")
        refus = client.post(f"/missions/{analyse}/modele", json={"modele": "BAIL_COMMERCIAL"})
        assert refus.status_code == 409
        assert client.get(f"/missions/{analyse}/livrable/word").status_code == 409
        client.cookies.clear()
        redaction = self._ouvrir(client)
        inconnu = client.post(f"/missions/{redaction}/modele", json={"modele": "VENTE_DE_LUNE"})
        assert inconnu.status_code == 404


class TestLaConsigneAuJuriste:
    def test_une_consigne_au_juriste_ne_part_pas_au_client(self):
        """18 octobre 2026 : « à vérifier par le juriste » s'adresse au juriste. Retenue
        telle quelle, la consigne partirait dans le contrat du client."""
        m = _redaction().ajouter(
            [
                {
                    "rubrique": "ARTICLE",
                    "titre": "Durée",
                    "texte": "Trois ans. Le renouvellement est à vérifier par le juriste.",
                }
            ],
            origine=OrigineProposition.MODELE,
        )
        retenue = m.decider(1, decision=DecisionProposition.RETENUE, par="C-007", a_l_instant=LE)
        with pytest.raises(TransitionMissionRefusee, match="consigne au juriste"):
            retenue.soumettre(par="C-007", a_l_instant=LE)
