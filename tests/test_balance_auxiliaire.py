"""Les balances auxiliaires clients et fournisseurs, annexes de la DSF (15 octobre 2026).

⚠️ CE QUI EST GARDÉ
* une ligne par tiers : reporté (journal AN), débit, crédit, solde ;
* les seules écritures validées ;
* ce qui n'a pas de tiers figure, sur une ligne à part, pour que le total retombe sur le
  compte collectif ;
* les clients (41) et les fournisseurs (40) ne se mêlent pas ;
* l'Excel porte des NOMBRES, le NIU du tiers, et dit qu'il n'est pas un modèle de la DGI ;
* un compte d'un autre dossier n'y accède pas.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.contextes.comptabilite.contrats import (
    EcritureComptable,
    EtatEcriture,
    LigneEcriture,
    Sens,
)
from app.contextes.comptabilite.domaine.balance_auxiliaire import (
    NatureAuxiliaire,
    balance_auxiliaire,
)
from app.contextes.transverse.api import NIU_DEMO
from tests.conftest import exige_postgresql, ouvrir_une_session

ALUCAM = "M012345678901A"
SOCATRAL = "M098765432109B"


def _ecriture(*lignes, numero, journal="OD", etat=EtatEcriture.VALIDEE) -> EcritureComptable:
    return EcritureComptable(
        journal=journal,
        exercice="2026",
        numero=numero,
        date_operation=date(2026, 6, 30),
        libelle="Essai",
        piece_justificative=f"PJ-{numero}",
        lignes=[
            LigneEcriture(
                compte=c,
                libelle=c,
                sens=Sens.DEBIT if m > 0 else Sens.CREDIT,
                montant=abs(Decimal(m)),
                tiers=t,
            )
            for c, m, t in lignes
        ],
        etat=etat,
        saisie_par="l.fotso",
        saisie_le=datetime(2026, 6, 30, 9, 0),
        validee_par="a.bouba" if etat is EtatEcriture.VALIDEE else None,
        validee_le=datetime(2026, 7, 1, 9, 0) if etat is EtatEcriture.VALIDEE else None,
    )


ECRITURES = [
    _ecriture(("401", -800_000, ALUCAM), ("471", 800_000, None), numero=1, journal="AN"),
    _ecriture(("602", 4_000_000, None), ("401", -4_000_000, ALUCAM), numero=2),
    _ecriture(("401", 1_000_000, ALUCAM), ("521", -1_000_000, None), numero=3),
    _ecriture(("602", 1_500_000, None), ("4011", -1_500_000, SOCATRAL), numero=4),
    _ecriture(("605", 200_000, None), ("401", -200_000, None), numero=5),
    _ecriture(("411", 2_000_000, "CLI-01"), ("701", -2_000_000, None), numero=6),
    _ecriture(
        ("602", 999_000, None), ("401", -999_000, ALUCAM), numero=7, etat=EtatEcriture.BROUILLON
    ),
]


class TestLeCalcul:
    def test_une_ligne_par_fournisseur_avec_son_report(self):
        lignes = {x.tiers: x for x in balance_auxiliaire(ECRITURES, NatureAuxiliaire.FOURNISSEURS)}

        alucam = lignes[ALUCAM]
        assert (alucam.report, alucam.debit, alucam.credit) == (
            Decimal(-800_000),
            Decimal(1_000_000),
            Decimal(4_000_000),
        )
        assert alucam.solde == Decimal(-3_800_000)
        assert lignes[SOCATRAL].solde == Decimal(-1_500_000)  # un sous-compte 4011 compte

    def test_ce_qui_n_a_pas_de_tiers_figure_en_dernier(self):
        lignes = balance_auxiliaire(ECRITURES, NatureAuxiliaire.FOURNISSEURS)

        assert lignes[-1].tiers is None and lignes[-1].solde == Decimal(-200_000)
        total = sum((x.solde for x in lignes), Decimal(0))
        assert total == Decimal(-800_000 - 4_000_000 + 1_000_000 - 1_500_000 - 200_000)

    def test_les_clients_ne_se_melent_pas_aux_fournisseurs(self):
        (client,) = balance_auxiliaire(ECRITURES, NatureAuxiliaire.CLIENTS)

        assert client.tiers == "CLI-01" and client.solde == Decimal(2_000_000)

    def test_un_brouillon_ne_compte_pas(self):
        lignes = {x.tiers: x for x in balance_auxiliaire(ECRITURES, NatureAuxiliaire.FOURNISSEURS)}

        assert lignes[ALUCAM].credit == Decimal(4_000_000)


# ── Par les routes ──────────────────────────────────────────────────────────────

BATIMENT = NIU_DEMO["BATIMENT"]
COMPTABLE = "l.fotso@cga-brcg.cm"
URL = f"/comptabilite/dossiers/{BATIMENT}/balances-auxiliaires/fournisseurs"


def _une_dette_validee(client, tiers: str, montant: int) -> None:
    saisie = client.post(
        f"/comptabilite/dossiers/{BATIMENT}/ecritures",
        json={
            "journal": "OD",
            "exercice": "2026",
            "date_operation": "2026-06-15",
            "libelle": "Facture fournisseur",
            "piece_justificative": f"FA-{tiers[-4:]}",
            "lignes": [
                {"compte": "602", "libelle": "Achat", "sens": "DEBIT", "montant": str(montant)},
                {
                    "compte": "401",
                    "libelle": "Fournisseur",
                    "sens": "CREDIT",
                    "montant": str(montant),
                    "tiers": tiers,
                },
            ],
        },
    )
    assert saisie.status_code == 201, saisie.text
    numero = saisie.json()["numero"]
    validation = client.post(
        f"/comptabilite/dossiers/{BATIMENT}/ecritures/2026/OD/{numero}/validation"
    )
    assert validation.status_code == 200, validation.text


@exige_postgresql
class TestParLesRoutes:
    def test_la_balance_et_son_excel(self, plateforme):
        client = plateforme
        ouvrir_une_session(client, COMPTABLE)
        _une_dette_validee(client, ALUCAM, 2_500_000)

        lignes = client.get(URL, params={"exercice": "2026"})
        assert lignes.status_code == 200, lignes.text
        (alucam,) = [x for x in lignes.json() if x["tiers"] == ALUCAM]
        assert Decimal(alucam["solde"]) == Decimal(-2_500_000)
        # Son code a la forme d'un NIU au référentiel : il est repris comme tel.
        assert alucam["niu"] == ALUCAM

        excel = client.get(f"{URL}/excel", params={"exercice": "2026"})
        assert excel.status_code == 200, excel.text
        classeur = load_workbook(BytesIO(excel.content))
        feuille = classeur["Balance"]
        rangees = list(feuille.iter_rows(values_only=True))
        assert rangees[0][0] == "Code tiers"
        (ligne,) = [r for r in rangees if r[0] == ALUCAM]
        assert ligne[1] == ALUCAM and ligne[7] == 2_500_000  # un NOMBRE, au crédit
        assert rangees[-1][0] == "Total"
        # Toutes les cellules de montant, total compris : une machine les relira.
        assert all(isinstance(v, int) for r in rangees[1:] for v in r[3:]), rangees
        lire = " ".join(str(r[0]) for r in classeur["À lire"].iter_rows(values_only=True))
        assert "Ce n'est pas un modèle de la DGI" in lire

    def test_un_autre_dossier_n_y_accede_pas(self, plateforme):
        client = plateforme
        ouvrir_une_session(client, "jp.nkoa@batimentplus.cm")
        autre = f"/comptabilite/dossiers/{NIU_DEMO['COLOMBE']}/balances-auxiliaires/fournisseurs"

        assert client.get(autre, params={"exercice": "2026"}).status_code == 403
        assert client.get(f"{autre}/excel", params={"exercice": "2026"}).status_code == 403
