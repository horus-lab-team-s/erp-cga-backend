"""Le registre des immobilisations et leur amortissement (17 octobre 2026).

─────────────────────────────────────────────────────────────────────────────────
CE QUE CE FICHIER GARDE

1. **Le calcul** : linéaire, prorata temporis en jours, base 365 ; le cumul plafonné à la
   valeur d'origine, la dernière dotation qui solde le bien ; la sortie qui arrête
   l'amortissement ; le terrain qui ne s'amortit pas.
2. **Le tableau** : brut, amortissements, valeur nette, par catégorie ; acquisitions et
   sorties de l'exercice ; un total qui est la somme des lignes.
3. **La fiche** : la base figée dès qu'une dotation est passée en écriture.
4. **Les routes** : saisir, sortir, lire le tableau, l'Excel ; la dotation en brouillon au
   journal OD, une seule vivante par exercice ; les accès ; PostgreSQL.
─────────────────────────────────────────────────────────────────────────────────
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from app.contextes.comptabilite.adaptateurs.sortant.depots_immobilisations import (
    vider_les_immobilisations,
)
from app.contextes.comptabilite.api import vider_les_ecritures_en_memoire
from app.contextes.comptabilite.domaine.immobilisations import (
    CategorieImmobilisation,
    Immobilisation,
    ImmobilisationRefusee,
    cumul_au,
    dotation_de_l_exercice,
    tableau_des_amortissements,
)
from app.contextes.transverse.api import MOT_DE_PASSE_DEMO, reinitialiser_atelier
from app.main import creer_application
from tests.conftest import exige_postgresql, ouvrir_une_session

BATIMENT = "M081234567890P"
COMPTABLE = "l.fotso@cga-brcg.cm"
COMPTABLE_AUTRE = "c.ndongo@cga-brcg.cm"
CHARGEE_DE_CLIENTELE = "p.moukouri@cga-brcg.cm"
LE = datetime(2026, 10, 17, 9)


@pytest.fixture(autouse=True)
def neuf():
    reinitialiser_atelier()
    vider_les_immobilisations()
    vider_les_ecritures_en_memoire()
    yield
    vider_les_immobilisations()
    vider_les_ecritures_en_memoire()


def _bien(**champs) -> Immobilisation:
    valeurs = dict(
        dossier=BATIMENT,
        code="INFO-01",
        libelle="Ordinateurs portables",
        categorie=CategorieImmobilisation.MATERIEL,
        valeur_origine=Decimal(3_650_000),
        mise_en_service=date(2026, 7, 1),
        duree_annees=4,
        enregistree_par="C-004",
        enregistree_le=LE,
    )
    return Immobilisation(**(valeurs | champs))


# ── Le calcul ─────────────────────────────────────────────────────────────────


class TestLeCalcul:
    def test_prorata_en_jours_la_premiere_annee(self):
        # 184 jours du 1er juillet au 31 décembre : 3 650 000 × 184 / (4 × 365).
        assert dotation_de_l_exercice(_bien(), date(2026, 1, 1), date(2026, 12, 31)) == 460_000

    def test_une_annee_pleine(self):
        assert dotation_de_l_exercice(_bien(), date(2027, 1, 1), date(2027, 12, 31)) == 912_500

    def test_le_cumul_ne_depasse_jamais_la_valeur_d_origine(self):
        bien = _bien()
        assert cumul_au(bien, date(2040, 1, 1)) == bien.valeur_origine
        total = sum(
            dotation_de_l_exercice(bien, date(a, 1, 1), date(a, 12, 31)) for a in range(2026, 2032)
        )
        assert total == bien.valeur_origine

    def test_la_derniere_dotation_solde_le_bien_malgre_les_arrondis(self):
        # Une valeur qui ne se divise pas : les annuités arrondies ne tombent pas juste, le
        # cumul plafonné, si.
        bien = _bien(
            valeur_origine=Decimal(1_000_003), duree_annees=3, mise_en_service=date(2026, 3, 17)
        )
        dotations = [
            dotation_de_l_exercice(bien, date(a, 1, 1), date(a, 12, 31)) for a in range(2026, 2031)
        ]
        assert sum(dotations) == Decimal(1_000_003)
        assert all(d >= 0 for d in dotations)

    def test_rien_avant_la_mise_en_service(self):
        assert dotation_de_l_exercice(_bien(), date(2025, 1, 1), date(2025, 12, 31)) == 0

    def test_la_sortie_arrete_l_amortissement(self):
        sorti = _bien(sortie_le=date(2027, 3, 31), motif_sortie="cession")
        # 90 jours en 2027 : 3 650 000 × 90 / 1 460.
        assert dotation_de_l_exercice(sorti, date(2027, 1, 1), date(2027, 12, 31)) == 225_000
        assert dotation_de_l_exercice(sorti, date(2028, 1, 1), date(2028, 12, 31)) == 0

    def test_un_terrain_ne_s_amortit_pas(self):
        terrain = _bien(code="TER-01", categorie=CategorieImmobilisation.TERRAIN, duree_annees=None)
        assert cumul_au(terrain, date(2030, 1, 1)) == 0


class TestLaFiche:
    def test_un_amortissable_a_une_duree(self):
        with pytest.raises(ValidationError, match="durée"):
            _bien(duree_annees=None)

    def test_un_terrain_n_a_pas_de_duree(self):
        with pytest.raises(ValidationError, match="terrain"):
            _bien(categorie=CategorieImmobilisation.TERRAIN, duree_annees=5)

    def test_une_sortie_dit_son_motif_et_suit_la_mise_en_service(self):
        with pytest.raises(ValidationError, match="motif"):
            _bien(sortie_le=date(2027, 1, 1))
        with pytest.raises(ValidationError, match="précède"):
            _bien(sortie_le=date(2026, 1, 1), motif_sortie="erreur")

    def test_la_base_se_fige_des_qu_une_dotation_est_passee(self):
        bien = _bien()
        assert bien.modifier(valeur_origine=Decimal(4_000_000)).valeur_origine == 4_000_000
        dote = bien.marquer_dotee("2026")
        with pytest.raises(ImmobilisationRefusee, match="valeur_origine"):
            dote.modifier(valeur_origine=Decimal(4_000_000))
        # Le libellé, lui, se corrige toujours.
        assert dote.modifier(libelle="Portables du cabinet").libelle == "Portables du cabinet"


class TestLeTableau:
    def test_brut_amortissements_et_valeur_nette_par_categorie(self):
        biens = [
            _bien(mise_en_service=date(2025, 1, 1)),
            _bien(
                code="INFO-02", mise_en_service=date(2026, 7, 1), valeur_origine=Decimal(730_000)
            ),
            _bien(
                code="VEH-01",
                libelle="Camionnette",
                mise_en_service=date(2024, 1, 1),
                sortie_le=date(2026, 6, 30),
                motif_sortie="cession",
            ),
            _bien(code="TER-01", categorie=CategorieImmobilisation.TERRAIN, duree_annees=None),
        ]
        t = tableau_des_amortissements(
            BATIMENT, "2026", date(2026, 1, 1), date(2026, 12, 31), biens
        )
        materiel = next(lg for lg in t.lignes if lg.compte == "24")
        assert materiel.brut_ouverture == 3_650_000 * 2
        assert materiel.acquisitions == 730_000
        assert materiel.sorties == 3_650_000
        assert materiel.brut_cloture == 3_650_000 + 730_000
        assert materiel.amortissements_cloture == (
            materiel.amortissements_ouverture + materiel.dotations - materiel.amortissements_sortis
        )
        assert (
            materiel.valeur_nette_cloture == materiel.brut_cloture - materiel.amortissements_cloture
        )
        # Le bien encore présent : 2 ans d'amortissement fin 2026.
        assert next(b for b in t.biens if b["code"] == "INFO-01")["cumul_cloture"] == "1825000"
        terrain = next(lg for lg in t.lignes if lg.compte == "22")
        assert terrain.dotations == 0 and terrain.valeur_nette_cloture == terrain.brut_cloture
        for champ in (
            "brut_cloture",
            "dotations",
            "amortissements_cloture",
            "valeur_nette_cloture",
        ):
            assert getattr(t.total, champ) == sum(getattr(lg, champ) for lg in t.lignes)

    def test_un_bien_sorti_avant_l_exercice_n_y_figure_pas(self):
        sorti = _bien(
            mise_en_service=date(2024, 1, 1), sortie_le=date(2025, 6, 30), motif_sortie="rebut"
        )
        t = tableau_des_amortissements(
            BATIMENT, "2026", date(2026, 1, 1), date(2026, 12, 31), [sorti]
        )
        assert t.lignes == [] and t.biens == []


# ── Les routes ────────────────────────────────────────────────────────────────


def _client(courriel):
    client = TestClient(creer_application())
    reponse = client.post(
        "/transverse/session", json={"courriel": courriel, "mot_de_passe": MOT_DE_PASSE_DEMO}
    )
    assert reponse.status_code == 200, reponse.text
    return client


REGISTRE = f"/comptabilite/dossiers/{BATIMENT}/immobilisations"
TABLEAU = f"/comptabilite/dossiers/{BATIMENT}/amortissements/2026"
FICHE = {
    "libelle": "Ordinateurs portables",
    "categorie": "24",
    "valeur_origine": "3650000",
    "mise_en_service": "2026-07-01",
    "duree_annees": 4,
    "fournisseur": "Quincaillerie du Wouri",
}


class TestLesRoutes:
    def test_saisir_lire_et_passer_la_dotation(self):
        client = _client(COMPTABLE)
        saisie = client.put(f"{REGISTRE}/info-01", json=FICHE)
        assert saisie.status_code == 200, saisie.text
        assert saisie.json()["code"] == "INFO-01"
        assert [i["code"] for i in client.get(REGISTRE).json()] == ["INFO-01"]
        tableau = client.get(TABLEAU).json()
        assert tableau["ecriture_de_dotation"] is None
        assert tableau["tableau"]["total"]["dotations"] == "460000"

        dotation = client.post(f"{TABLEAU}/dotation")
        assert dotation.status_code == 201, dotation.text
        ecriture = dotation.json()
        assert ecriture["journal"] == "OD" and ecriture["etat"] == "BROUILLON"
        assert ecriture["date_operation"] == "2026-12-31"
        assert [(lg["compte"], lg["sens"], lg["montant"]) for lg in ecriture["lignes"]] == [
            ("681", "DEBIT", "460000"),
            ("28", "CREDIT", "460000"),
        ]
        cle = f"2026/OD/{ecriture['numero']:06d}"
        assert client.get(TABLEAU).json()["ecriture_de_dotation"] == cle

        # Une seule dotation vivante : le double clic ne double pas la charge.
        double = client.post(f"{TABLEAU}/dotation")
        assert double.status_code == 409 and "déjà passée" in double.json()["detail"]
        # Et la base est désormais figée.
        figee = client.put(f"{REGISTRE}/INFO-01", json=FICHE | {"valeur_origine": "4000000"})
        assert figee.status_code == 409, figee.text

    def test_rien_a_doter(self):
        client = _client(COMPTABLE)
        assert client.post(f"{TABLEAU}/dotation").status_code == 409

    def test_un_exercice_inconnu(self):
        client = _client(COMPTABLE)
        assert (
            client.get(f"/comptabilite/dossiers/{BATIMENT}/amortissements/1999").status_code == 404
        )

    def test_sortir_un_bien(self):
        client = _client(COMPTABLE)
        client.put(f"{REGISTRE}/INFO-01", json=FICHE)
        sortie = client.post(
            f"{REGISTRE}/INFO-01/sortie", json={"sortie_le": "2026-09-30", "motif": "cession"}
        )
        assert sortie.status_code == 200, sortie.text
        assert (
            client.post(
                f"{REGISTRE}/INFO-01/sortie", json={"sortie_le": "2026-10-30", "motif": "cession"}
            ).status_code
            == 409
        )
        assert client.put(f"{REGISTRE}/INFO-01", json=FICHE).status_code == 409
        ligne = client.get(TABLEAU).json()["tableau"]["lignes"][0]
        assert ligne["sorties"] == "3650000" and ligne["brut_cloture"] == "0"

    def test_une_fiche_incoherente(self):
        client = _client(COMPTABLE)
        terrain = client.put(f"{REGISTRE}/TER-01", json=FICHE | {"categorie": "22"})
        assert terrain.status_code == 422 and "terrain" in terrain.json()["detail"]

    def test_l_excel_du_tableau(self):
        client = _client(COMPTABLE)
        client.put(f"{REGISTRE}/INFO-01", json=FICHE | {"libelle": '=HYPERLINK("x")'})
        reponse = client.get(f"{TABLEAU}/excel")
        assert reponse.status_code == 200
        classeur = load_workbook(BytesIO(reponse.content))
        assert classeur.sheetnames == ["Tableau", "Biens", "À lire"]
        assert classeur["Tableau"]["H2"].value == 460000
        # Un libellé qui ressemble à une formule reste du texte.
        assert classeur["Biens"]["B2"].data_type == "s"

    def test_les_acces(self):
        chargee = _client(CHARGEE_DE_CLIENTELE)
        assert chargee.put(f"{REGISTRE}/INFO-01", json=FICHE).status_code == 403
        assert chargee.post(f"{TABLEAU}/dotation").status_code == 403

    def test_la_dotation_se_valide_comme_toute_ecriture_et_se_journalise(self):
        client = _client(COMPTABLE)
        client.put(f"{REGISTRE}/INFO-01", json=FICHE)
        ecriture = client.post(f"{TABLEAU}/dotation").json()
        chemin = (
            f"/comptabilite/dossiers/{BATIMENT}/ecritures/2026/OD/{ecriture['numero']}/validation"
        )
        validee = client.post(chemin)
        assert validee.status_code == 200, validee.text
        assert validee.json()["etat"] == "VALIDEE"
        # Un comptable d'un autre portefeuille ne voit pas le dossier.
        assert _client(COMPTABLE_AUTRE).get(REGISTRE).status_code in (403, 404)


class TestUneDotationContrePassee:
    def test_libere_l_exercice_pour_une_nouvelle_dotation(self):
        client = _client(COMPTABLE)
        client.put(f"{REGISTRE}/INFO-01", json=FICHE)
        ecriture = client.post(f"{TABLEAU}/dotation").json()
        base = f"/comptabilite/dossiers/{BATIMENT}/ecritures/2026/OD/{ecriture['numero']}"
        assert client.post(f"{base}/validation").status_code == 200
        annulee = client.post(
            f"{base}/contre-passation",
            json={"motif": "durée revue avec le fiscaliste", "date_operation": "2026-12-31"},
        )
        assert annulee.status_code == 201, annulee.text
        assert client.get(TABLEAU).json()["ecriture_de_dotation"] is None
        nouvelle = client.post(f"{TABLEAU}/dotation")
        assert nouvelle.status_code == 201, nouvelle.text
        assert nouvelle.json()["numero"] > annulee.json()["numero"]


class TestLeControleAvecLaBalance:
    def test_un_bien_oublie_au_registre_se_voit(self):
        client = _client(COMPTABLE)
        ecritures = f"/comptabilite/dossiers/{BATIMENT}/ecritures"
        achat = client.post(
            ecritures,
            json={
                "journal": "OD",
                "exercice": "2026",
                "date_operation": "2026-07-01",
                "libelle": "Achat de portables",
                "piece_justificative": "F-INFO-01",
                "lignes": [
                    {"compte": "24", "libelle": "Portables", "sens": "DEBIT", "montant": "3650000"},
                    {
                        "compte": "401",
                        "libelle": "Portables",
                        "sens": "CREDIT",
                        "montant": "3650000",
                    },
                ],
            },
        ).json()
        client.post(f"{ecritures}/2026/OD/{achat['numero']}/validation")
        controle = client.get(TABLEAU).json()["controle"]
        assert controle["brut_balance"] == "3650000" and controle["brut_registre"] == "0"
        assert controle["concorde"] is False

        client.put(f"{REGISTRE}/INFO-01", json=FICHE)
        dotation = client.post(f"{TABLEAU}/dotation").json()
        # La dotation en brouillon n'est pas encore dans la balance : l'écart le dit.
        assert client.get(TABLEAU).json()["controle"]["concorde"] is False
        client.post(f"{ecritures}/2026/OD/{dotation['numero']}/validation")
        controle = client.get(TABLEAU).json()["controle"]
        assert controle["amortissements_balance"] == "460000"
        assert controle["concorde"] is True


class TestLaPieceDeLaDsf:
    def test_le_tableau_est_une_piece_produite_de_la_dsf_normale(self):
        from app.contextes.cloture.domaine.entites import SystemeDsf
        from app.contextes.cloture.domaine.pieces_dsf import EtatPiece, etat_des_pieces

        sept = next(lg for lg in etat_des_pieces(SystemeDsf.NORMAL, {}) if lg.piece.numero == 7)
        assert sept.etat is EtatPiece.PRODUITE
        assert sept.piece.produite == "tableau-amortissements"


class TestSurPostgresql:
    pytestmark = exige_postgresql

    def test_le_registre_et_la_dotation_d_une_requete_a_l_autre(self, plateforme):
        client = plateforme
        ouvrir_une_session(client, COMPTABLE)
        assert client.put(f"{REGISTRE}/INFO-01", json=FICHE).status_code == 200
        assert (
            client.put(f"{REGISTRE}/INFO-01", json=FICHE | {"libelle": "Portables"}).status_code
            == 200
        )
        assert [i["libelle"] for i in client.get(REGISTRE).json()] == ["Portables"]
        assert client.post(f"{TABLEAU}/dotation").status_code == 201
        assert client.get(REGISTRE).json()[0]["dotations_passees"] == ["2026"]
        assert client.post(f"{TABLEAU}/dotation").status_code == 409
