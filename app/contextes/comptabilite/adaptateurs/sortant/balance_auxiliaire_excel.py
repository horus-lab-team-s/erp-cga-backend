"""La balance auxiliaire en Excel, comme la DGI la demande en annexe de la DSF (15 octobre 2026).

⚠️ PAS ENCORE LE MODÈLE DGI : celui-ci ne se télécharge qu'après connexion à la
télédéclaration. La feuille « Balance » porte le contenu, sans fioriture, pour se
recopier dans le modèle ou s'y aligner dès qu'il sera obtenu ; la feuille « À lire » le
dit, pour que personne ne téléverse ce fichier tel quel en croyant au format officiel.

⚠️ LES MONTANTS SONT DES NOMBRES, pas du texte mis en forme : un fichier Excel destiné à
être relu par une machine ne doit pas porter « 1 500 000 » en chaîne.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.contextes.comptabilite.domaine.balance_auxiliaire import (
    LigneBalanceAuxiliaire,
    NatureAuxiliaire,
)
from app.partage.horloge import a_douala

__all__ = ["EnteteBalanceAuxiliaire", "FicheDuTiers", "rendre_la_balance_auxiliaire"]

_TITRES = {
    NatureAuxiliaire.CLIENTS: "Balance auxiliaire clients",
    NatureAuxiliaire.FOURNISSEURS: "Balance auxiliaire fournisseurs",
}
_COLONNES = (
    "Code tiers",
    "NIU",
    "Raison sociale",
    "Solde reporté",
    "Mouvements débit",
    "Mouvements crédit",
    "Solde débiteur",
    "Solde créditeur",
)
_FORMAT = "#,##0;-#,##0;0"


@dataclass(frozen=True)
class EnteteBalanceAuxiliaire:
    niu: str
    denomination: str
    exercice: str
    nature: NatureAuxiliaire
    editee_le: datetime


@dataclass(frozen=True)
class FicheDuTiers:
    denomination: str
    niu: str | None


def rendre_la_balance_auxiliaire(
    lignes: list[LigneBalanceAuxiliaire],
    entete: EnteteBalanceAuxiliaire,
    fiches: dict[str, FicheDuTiers],
) -> bytes:
    """Le classeur, en octets. `fiches` : par code de tiers, sa raison sociale et son NIU,
    établis par l'appelant (registre du dossier, format du NIU au référentiel). Un tiers
    sans fiche garde son code, sans raison sociale ni NIU."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Balance"
    feuille.append(_COLONNES)
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
        cellule.fill = PatternFill("solid", fgColor="EEEAF5")
        cellule.alignment = Alignment(wrap_text=True, vertical="top")

    # ⚠️ Le total additionne les montants AFFICHÉS : arrondis ligne à ligne, des montants
    # exacts sommés donneraient un total que la colonne ne fait pas.
    totaux = [0] * 5
    for ligne in lignes:
        fiche = fiches.get(ligne.tiers or "")
        solde = ligne.solde
        debiteur = solde if solde > 0 else Decimal(0)
        crediteur = -solde if solde < 0 else Decimal(0)
        niu = fiche.niu if fiche else None
        feuille.append(
            [
                ligne.tiers or "(sans tiers)",
                niu or "",
                fiche.denomination if fiche else ("" if ligne.tiers else "À ventiler par tiers"),
                _entier(ligne.report),
                _entier(ligne.debit),
                _entier(ligne.credit),
                _entier(debiteur),
                _entier(crediteur),
            ]
        )
        for i, v in enumerate((ligne.report, ligne.debit, ligne.credit, debiteur, crediteur)):
            totaux[i] += _entier(v)
    feuille.append(["Total", "", "", *totaux])
    for cellule in feuille[feuille.max_row]:
        cellule.font = Font(bold=True)
    for rang in range(2, feuille.max_row + 1):
        # ⚠️ Un code ou une raison sociale qui commence par « = » deviendrait une formule
        # dans le tableur de qui ouvre le fichier : les trois colonnes de texte restent du
        # texte, quoi qu'elles contiennent.
        for colonne in range(1, 4):
            feuille.cell(row=rang, column=colonne).data_type = "s"
        for colonne in range(4, 9):
            feuille.cell(row=rang, column=colonne).number_format = _FORMAT
    for lettre, largeur in zip("ABCDEFGH", (18, 18, 40, 16, 16, 16, 16, 16), strict=True):
        feuille.column_dimensions[lettre].width = largeur
    feuille.freeze_panes = "A2"

    lire = classeur.create_sheet("À lire")
    for texte in (
        _TITRES[entete.nature],
        f"{entete.denomination} · NIU {entete.niu} · exercice {entete.exercice}",
        f"Éditée le {a_douala(entete.editee_le):%d/%m/%Y à %H:%M} (heure de Douala), "
        "sur les écritures validées.",
        "",
        "Solde reporté : journal des à-nouveaux. Mouvements : l'exercice, hors à-nouveaux.",
        "Une ligne « (sans tiers) » signale des écritures à ventiler avant le dépôt.",
        "",
        "Ce fichier porte le CONTENU de la balance auxiliaire demandée en annexe de la DSF.",
        "Il n'est pas encore au modèle Excel de la DGI, qui ne se télécharge qu'après",
        "connexion à la télédéclaration : recopier dans ce modèle avant de téléverser.",
    ):
        lire.append([texte])
    lire["A1"].font = Font(bold=True, size=13)
    lire.column_dimensions["A"].width = 100

    tampon = BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


def _entier(valeur: Decimal) -> int:
    """Au franc, arrondi commercial : `int()` tronquait, et le total calculé sur les
    montants exacts pouvait différer de la somme des lignes affichées."""
    return int(valeur.quantize(Decimal(1), rounding=ROUND_HALF_UP))
