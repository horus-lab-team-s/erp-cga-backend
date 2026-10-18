"""Le tableau des amortissements de l'exercice en Excel, annexe de la DSF (17 octobre 2026).

⚠️ PAS UN MODÈLE DGI : la feuille porte le contenu du tableau (valeurs brutes,
amortissements, valeurs nettes, par catégorie, puis bien par bien), et la feuille « À lire »
le dit. Les montants sont des nombres ; les textes restent du texte, jamais une formule.
"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.contextes.comptabilite.domaine.immobilisations import TableauDesAmortissements
from app.partage.horloge import a_douala

__all__ = ["rendre_le_tableau_des_amortissements"]

_FORMAT = "#,##0;-#,##0;0"
_ENTETE = PatternFill("solid", fgColor="EEEAF5")
_COLONNES = (
    "Compte",
    "Catégorie",
    "Brut à l'ouverture",
    "Acquisitions",
    "Sorties",
    "Brut à la clôture",
    "Amortissements à l'ouverture",
    "Dotations de l'exercice",
    "Amortissements des sorties",
    "Amortissements à la clôture",
    "Valeur nette à la clôture",
)
_CHAMPS = (
    "brut_ouverture",
    "acquisitions",
    "sorties",
    "brut_cloture",
    "amortissements_ouverture",
    "dotations",
    "amortissements_sortis",
    "amortissements_cloture",
    "valeur_nette_cloture",
)


def _entete(feuille, colonnes) -> None:
    feuille.append(colonnes)
    for cellule in feuille[feuille.max_row]:
        cellule.font = Font(bold=True)
        cellule.fill = _ENTETE
        cellule.alignment = Alignment(wrap_text=True, vertical="top")


def _texte_et_nombres(feuille, rang: int, textes: int, jusqu_a: int) -> None:
    for colonne in range(1, textes + 1):
        feuille.cell(row=rang, column=colonne).data_type = "s"
    for colonne in range(textes + 1, jusqu_a + 1):
        feuille.cell(row=rang, column=colonne).number_format = _FORMAT


def rendre_le_tableau_des_amortissements(
    tableau: TableauDesAmortissements, *, denomination: str, editee_le: datetime
) -> bytes:
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Tableau"
    _entete(feuille, _COLONNES)
    for ligne in [*tableau.lignes, tableau.total]:
        feuille.append([ligne.compte, ligne.libelle, *(int(getattr(ligne, c)) for c in _CHAMPS)])
        _texte_et_nombres(feuille, feuille.max_row, 2, len(_COLONNES))
    for cellule in feuille[feuille.max_row]:
        cellule.font = Font(bold=True)
    for i, largeur in enumerate((10, 44, *([16] * 9))):
        feuille.column_dimensions[chr(ord("A") + i)].width = largeur
    feuille.freeze_panes = "C2"

    biens = classeur.create_sheet("Biens")
    _entete(
        biens,
        (
            "Code",
            "Libellé",
            "Compte",
            "Mise en service",
            "Durée (ans)",
            "Sortie le",
            "Valeur d'origine",
            "Dotation de l'exercice",
            "Amortissements cumulés",
            "Valeur nette",
        ),
    )
    for b in tableau.biens:
        biens.append(
            [
                b["code"],
                b["libelle"],
                b["compte"],
                b["mise_en_service"],
                str(b["duree_annees"] or ""),
                b["sortie_le"] or "",
                int(Decimal(b["valeur_origine"])),
                int(Decimal(b["dotation"])),
                int(Decimal(b["cumul_cloture"])),
                int(Decimal(b["valeur_nette_cloture"])),
            ]
        )
        _texte_et_nombres(biens, biens.max_row, 6, 10)
    for i, largeur in enumerate((14, 40, 8, 14, 10, 12, 16, 16, 16, 16)):
        biens.column_dimensions[chr(ord("A") + i)].width = largeur
    biens.freeze_panes = "A2"

    lire = classeur.create_sheet("À lire")
    for texte in (
        "Tableau des amortissements",
        f"{denomination} · NIU {tableau.dossier} · exercice {tableau.exercice} "
        f"({tableau.ouverture:%d/%m/%Y} au {tableau.cloture:%d/%m/%Y})",
        f"Édité le {a_douala(editee_le):%d/%m/%Y à %H:%M} (heure de Douala), "
        "depuis le registre des immobilisations.",
        "",
        "Amortissement linéaire, prorata temporis en jours, base de 365 jours par année.",
        "Les durées sont saisies bien par bien ; elles se valident avec le fiscaliste.",
        "",
        "Ce fichier porte le CONTENU du tableau demandé en annexe de la DSF.",
        "Ce n'est pas un modèle de la DGI.",
    ):
        lire.append([texte])
    lire["A1"].font = Font(bold=True, size=13)
    lire.column_dimensions["A"].width = 100

    tampon = BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()
