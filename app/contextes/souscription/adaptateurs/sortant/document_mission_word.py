"""Le projet de contrat en Word (18 octobre 2026, point 23).

─────────────────────────────────────────────────────────────────────────────────
Le PDF du livrable se lit et se valide ; il ne se retouche pas. Un contrat, si : les
parties le relisent, le juriste ajuste un délai en réunion, le client y porte ses
initiales. Le même projet part donc aussi en `.docx`, fait des mêmes articles retenus,
dans le même ordre et avec la même numérotation que le PDF (`_projet_de_contrat`).

⚠️ TANT QUE LA MISSION N'EST PAS LIVRÉE, le document le dit en tête : « PROJET · non
validé ». Comme le PDF.

⚠️ RIEN D'INTERNE (point 21) : ni le nom du juriste, ni celui du valideur.
─────────────────────────────────────────────────────────────────────────────────
"""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.contextes.souscription.domaine.missions import EtatMission, Mission, Rubrique

__all__ = ["TYPE_WORD", "rendre_le_contrat_en_word"]

TYPE_WORD = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# La palette des documents du cabinet (voir `document_proforma.py`).
_INDIGO = RGBColor(0x2E, 0x1B, 0x4D)
_GRIS = RGBColor(0x5B, 0x55, 0x6B)
_MAGENTA = RGBColor(0xB0, 0x1E, 0x6B)


def rendre_le_contrat_en_word(mission: Mission) -> bytes:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    document.core_properties.author = "CGA Broad Range Consulting Group"
    document.core_properties.title = mission.type_de_contrat or "Projet de contrat"

    section = document.sections[0]
    entete = section.header.paragraphs[0]
    marque = entete.add_run("CGA Broad Range Consulting Group")
    marque.bold, marque.font.size, marque.font.color.rgb = True, Pt(9), _INDIGO
    if mission.etat is not EtatMission.LIVREE:
        projet = entete.add_run("    PROJET · non validé")
        projet.bold, projet.font.size, projet.font.color.rgb = True, Pt(9), _MAGENTA
    pied = section.footer.paragraphs[0].add_run(
        "CGA Broad Range Consulting Group · agrément MINFI/DGI n° 00000048"
    )
    pied.font.size, pied.font.color.rgb = Pt(8), _GRIS

    titre = document.add_paragraph()
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    texte = titre.add_run(f"PROJET DE {(mission.type_de_contrat or 'contrat').upper()}")
    texte.bold, texte.font.size, texte.font.color.rgb = True, Pt(15), _INDIGO

    note = document.add_paragraph().add_run(
        "Projet préparé par le cabinet, à relire par les parties avant toute signature."
    )
    note.italic, note.font.size, note.font.color.rgb = True, Pt(9.5), _GRIS

    articles = [p for p in mission.retenues if p.rubrique is Rubrique.ARTICLE]
    for numero, article in enumerate(articles, start=1):
        tete = document.add_paragraph()
        tete.paragraph_format.space_before = Pt(10)
        tete.paragraph_format.keep_with_next = True
        intitule = tete.add_run(f"Article {numero} · {article.titre}")
        intitule.bold, intitule.font.color.rgb = True, _INDIGO
        for alinea in (article.texte_retenu or "").split("\n"):
            if alinea.strip():
                corps = document.add_paragraph(alinea.strip())
                corps.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    tampon = BytesIO()
    document.save(tampon)
    return tampon.getvalue()
